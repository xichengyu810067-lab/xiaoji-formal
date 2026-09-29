"""Rebuild the exact-headword word-chain index from the official MOE archive.

Install the pinned build dependency in scripts/requirements-word-chain.txt, then run:
  python scripts/build-word-chain-lexicon.py <official-archive.zip>

The source archive is never copied into the application package. The index contains
only unmodified headwords. Keep the official usage PDF beside the generated index.
"""

import hashlib
import io
import json
import pathlib
import sys
import unicodedata
import zipfile

from openpyxl import load_workbook


VERSION = "2015_20260625"
ARCHIVE_SHA256 = "64003a98fcc7097940e5a536c999bc08ba7c07e2c1be66448f01bf1ae10a53fc"
INNER_XLSX = f"dict_revised_{VERSION}.xlsx"
SOURCE_URL = (
    "https://language.moe.gov.tw/001/Upload/Files/site_content/M0001/"
    f"respub/download/dict_revised_{VERSION}.zip"
)
USAGE_URL = (
    "https://language.moe.gov.tw/001/Upload/Files/site_content/"
    "M0001/respub/reviseddict_10312.pdf"
)
OUT_DIR = pathlib.Path(__file__).resolve().parents[1] / "assets" / "word-chain"


def sha256(blob):
    return hashlib.sha256(blob).hexdigest()


def is_han_headword(word):
    return (
        isinstance(word, str)
        and 2 <= len(word) <= 6
        and all(
            unicodedata.name(char, "").startswith(("CJK UNIFIED IDEOGRAPH", "CJK COMPATIBILITY IDEOGRAPH"))
            for char in word
        )
    )


def build(archive_path):
    archive = archive_path.read_bytes()
    digest = sha256(archive)
    if digest != ARCHIVE_SHA256:
        raise ValueError(f"Expected official {VERSION} archive SHA-256 {ARCHIVE_SHA256}, got {digest}")

    with zipfile.ZipFile(io.BytesIO(archive)) as source:
        workbook_bytes = source.read(INNER_XLSX)
    workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        if header[0] != "字詞名":
            raise ValueError("Official headword column changed; inspect the new source before rebuilding")
        all_headwords = set()
        eligible_rows = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            word = row[0]
            if is_han_headword(word):
                eligible_rows += 1
                all_headwords.add(word)
    finally:
        workbook.close()

    words = sorted(all_headwords)
    output = ("\n".join(words) + "\n").encode("utf-8")
    usage_path = OUT_DIR / "MOE-usage-revised.pdf"
    if not usage_path.is_file():
        raise FileNotFoundError(f"Official complete usage instructions are required: {usage_path}")
    metadata = {
        "source": "中華民國教育部《重編國語辭典修訂本》",
        "sourceVersion": VERSION,
        "sourceUrl": SOURCE_URL,
        "sourceArchiveSha256": digest,
        "sourceXlsxSha256": sha256(workbook_bytes),
        "usageInstructionsUrl": USAGE_URL,
        "usageInstructionsFile": usage_path.name,
        "usageInstructionsSha256": sha256(usage_path.read_bytes()),
        "license": "CC BY-ND 3.0 TW",
        "eligibleSourceRows": eligible_rows,
        "uniqueHeadwords": len(words),
        "indexFile": f"moe-revised-{VERSION}.txt",
        "indexSha256": sha256(output),
    }
    (OUT_DIR / metadata["indexFile"]).write_bytes(output)
    (OUT_DIR / "moe-revised-source.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n"
    )
    print(f"{len(words)} original headwords; index SHA-256 {metadata['indexSha256']}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: python scripts/build-word-chain-lexicon.py <official-archive.zip>")
    build(pathlib.Path(sys.argv[1]))
